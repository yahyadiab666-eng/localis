#!/usr/bin/env python3
"""Importación de Excel (.xlsx) con varias pestañas.

Verifica que el importador elija la hoja de inventario y no una pestaña de
resumen/estadísticas, que antes rompía la detección de cabeceras con el error
«No pudimos reconocer una fila de encabezados con columnas de producto…».

Casos:
  1. Hoja activa = resumen, inventario en la 2ª pestaña -> se elige inventario.
  2. Inventario en la 1ª pestaña, activa = resumen -> se elige inventario.
  3. ERP con preámbulo en una sola hoja -> sigue funcionando.
  4. Solo hojas de resumen -> error claro (sin excepción).
"""

from __future__ import annotations

import io
import sys
from pathlib import Path
from types import SimpleNamespace

RAIZ = Path(__file__).resolve().parents[1]
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

_ERRORES = []


def _ok(condicion, mensaje):
    if condicion:
        print(f'  OK  {mensaje}')
        return True
    print(f'  FALLO  {mensaje}')
    _ERRORES.append(mensaje)
    return False


def _archivo(nombre, data):
    return SimpleNamespace(filename=nombre, stream=io.BytesIO(data))


def _build(sheets, active=None):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        for fila in rows:
            ws.append(fila)
    if active:
        wb.active = wb.sheetnames.index(active)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _leer(data):
    from backend.inventory_import import (
        analizar_inventario,
        cargar_archivo_inventario,
        iter_lotes_productos,
        leer_encabezados_inventario,
    )

    cargado, extension, error = cargar_archivo_inventario(_archivo('inv.xlsx', data))
    if error:
        return [], error
    indice, encabezados, mapeo, error = analizar_inventario(cargado, extension)
    if error:
        return [], error
    productos = []
    for lote in iter_lotes_productos(
        cargado,
        extension,
        encabezados,
        {campo: campo for campo in mapeo},
        {'precio_en_bs': False},
        columnas=mapeo,
        fila_inicio=indice,
    ):
        productos.extend(lote)
    return productos, None


def main() -> int:
    print('\n=== Excel .xlsx multi-pestaña ===')

    data = _build(
        [
            ('Estadisticas', [['Resumen'], ['Total productos', 123], ['Ventas', 500]]),
            (
                'Inventario',
                [
                    ['codigo_barras', 'nombre', 'precio', 'stock'],
                    ['7591234567890', 'Licuadora Oster', 45.0, 3],
                    ['7591234567891', 'Teléfono Samsung A15', 180.0, 2],
                ],
            ),
        ],
        active='Estadisticas',
    )
    productos, error = _leer(data)
    _ok(error is None, f'hoja activa = resumen: no falla ({error})')
    _ok(
        [p['nombre'] for p in productos] == ['Licuadora Oster', 'Teléfono Samsung A15'],
        f'extrae la pestaña "Inventario" ({[p["nombre"] for p in productos]})',
    )

    data2 = _build(
        [
            ('Inventario', [['codigo', 'nombre', 'precio'], ['1', 'Radio', 20.0]]),
            ('Resumen', [['Resumen'], ['x', 1]]),
        ],
        active='Resumen',
    )
    productos2, error2 = _leer(data2)
    _ok(error2 is None and [p['nombre'] for p in productos2] == ['Radio'], 'usa la primera pestaña de producto')

    data3 = _build(
        [
            (
                'Hoja1',
                [
                    ['Reporte ERP'],
                    ['Generado', '2026'],
                    ['codigo', 'descripcion', 'costo'],
                    ['2', 'Ventilador', 15.0],
                ],
            )
        ]
    )
    productos3, error3 = _leer(data3)
    _ok(error3 is None and [p['nombre'] for p in productos3] == ['Ventilador'], 'sigue soportando ERP con preámbulo')

    data4 = _build([('Resumen', [['Totales'], ['a', 1]]), ('Gráficos', [['x'], [1]])])
    productos4, error4 = _leer(data4)
    _ok(not productos4 and error4 and 'encabezados' in error4, 'solo resúmenes -> error claro y controlado')

    print('\n=== RESULTADO ===')
    if _ERRORES:
        for item in _ERRORES:
            print(f'  - {item}')
        print(f'FALLOS: {len(_ERRORES)}')
        return 1
    print('OK importación .xlsx: elige la hoja de inventario e ignora pestañas de resumen')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
