#!/usr/bin/env python3
"""Importación .csv y .xlsx de extremo a extremo (lector + sincronización espejo).

Ejercita la misma función que usa la ruta de importación
(``backend.stores.procesar_csv_productos``) y valida:

  1. El lector .xlsx elige la hoja de inventario e ignora una hoja de resumen.
  2. Tras importar, los productos ausentes del archivo se ELIMINAN de la BD
     (LOCALIS_CSV_BAJAS=eliminar), dejando el inventario igual al archivo.
  3. Un comercio EN EL LÍMITE de su plan puede SUSTITUIR su catálogo: los
     ausentes se borran y liberan cupo (antes se contaba "actuales + altas" y se
     bloqueaba la importación).
  4. No hay excepciones no controladas en ninguno de los dos formatos.

Crea un comercio temporal y lo elimina al terminar.
"""

from __future__ import annotations

import io
import os
import sys
from pathlib import Path
from types import SimpleNamespace

os.environ.setdefault('LOCALIS_CSV_BAJAS', 'eliminar')
os.environ.setdefault('LOCALIS_IMG_PIPELINE', '0')  # sin red en la prueba

RAIZ = Path(__file__).resolve().parents[1]
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

try:
    from dotenv import load_dotenv

    load_dotenv(RAIZ / '.env', override=False)
except Exception:
    pass

_ERRORES = []
CORREO = '__fmt_test__@localis.test'
NOMBRE_COMERCIO = '__fmt_test__ comercio'


def _ok(condicion, mensaje):
    if condicion:
        print(f'  OK  {mensaje}')
        return True
    print(f'  FALLO  {mensaje}')
    _ERRORES.append(mensaje)
    return False


def _csv(productos):
    lineas = ['codigo,nombre,categoria,precio,existencia']
    for codigo, nombre, categoria, precio, stock in productos:
        lineas.append(f'{codigo},{nombre},{categoria},{precio},{stock}')
    return ('\n'.join(lineas) + '\n').encode('utf-8')


def _xlsx(productos, resumen_activo=True):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    resumen = wb.create_sheet('Resumen')
    resumen.append(['Reporte de inventario'])
    resumen.append(['Total productos', len(productos)])
    hoja = wb.create_sheet('Inventario')
    hoja.append(['Codigo', 'Nombre', 'Categoria', 'Precio', 'Existencia'])
    for fila in productos:
        hoja.append(list(fila))
    if resumen_activo:
        wb.active = wb.sheetnames.index('Resumen')
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _importar(comercio_id, nombre, contenido):
    """Igual que la ruta: procesa el archivo en el importador real."""
    from backend.stores import procesar_csv_productos

    archivo = SimpleNamespace(filename=nombre, stream=io.BytesIO(contenido))
    return procesar_csv_productos(comercio_id, archivo)


def _codigos(comercio_id):
    from backend.db import get_db_connection

    with get_db_connection(row_factory=True) as conexion:
        cursor = conexion.cursor()
        cursor.execute(
            'SELECT codigo_barras FROM productos WHERE comercio_id = ?',
            (comercio_id,),
        )
        return {dict(f)['codigo_barras'] for f in cursor.fetchall()}


def _limpiar(get_db_connection):
    with get_db_connection() as conexion:
        cursor = conexion.cursor()
        cursor.execute('SELECT id FROM comercios WHERE nombre = ?', (NOMBRE_COMERCIO,))
        fila = cursor.fetchone()
        if fila:
            cid = fila[0] if not isinstance(fila, dict) else list(fila.values())[0]
            cursor.execute('DELETE FROM productos WHERE comercio_id = ?', (cid,))
            cursor.execute('DELETE FROM comercios WHERE id = ?', (cid,))
        cursor.execute('DELETE FROM usuarios WHERE correo = ?', (CORREO,))
        conexion.commit()


def main() -> int:
    from backend.db import get_db_connection
    from backend.inventory_import import _modo_bajas

    print('\n=== Importación .csv / .xlsx de extremo a extremo ===')
    _ok(_modo_bajas() == 'eliminar', f'modo de bajas = eliminar ({_modo_bajas()})')

    _limpiar(get_db_connection)
    with get_db_connection() as conexion:
        cursor = conexion.cursor()
        cursor.execute(
            "INSERT INTO usuarios (nombre, correo, rol) VALUES (?, ?, 'comerciante') RETURNING id",
            ('Fmt Test', CORREO),
        )
        fila = cursor.fetchone()
        usuario_id = fila[0] if not isinstance(fila, dict) else list(fila.values())[0]
        cursor.execute(
            """
            INSERT INTO comercios (usuario_id, nombre, visible, estado_pago, limite_productos)
            VALUES (?, ?, 0, 'gratis', 3) RETURNING id
            """,
            (usuario_id, NOMBRE_COMERCIO),
        )
        fila = cursor.fetchone()
        comercio_id = fila[0] if not isinstance(fila, dict) else list(fila.values())[0]
        conexion.commit()

    base = [
        ('7701000000001', 'Televisor LG 50', 'Hogar', 250.0, 3),
        ('7701000000002', 'Licuadora Oster', 'Hogar', 45.0, 5),
        ('7701000000003', 'Telefono Samsung A15', 'Tecnologia', 180.0, 2),
    ]
    reemplazo = [
        ('7702000000001', 'Nevera Samsung 400L', 'Hogar', 600.0, 1),
        ('7702000000002', 'Microondas LG', 'Hogar', 120.0, 4),
        ('7702000000003', 'Ventilador Oster', 'Hogar', 30.0, 7),
    ]

    try:
        exito, mensaje, meta = _importar(comercio_id, 'inventario.csv', _csv(base))
        _ok(bool(exito), f'.csv inicial importa sin error ({mensaje[:60]})')
        _ok(meta.get('insertados') == 3, f'.csv: 3 altas ({meta.get("insertados")})')
        _ok(_codigos(comercio_id) == {p[0] for p in base}, '.csv deja los 3 productos en la BD')

        # El comercio está en el tope de su plan (3). Sustituir el catálogo por
        # otros 3 distintos debe funcionar: la sincronización espejo borra los
        # ausentes y libera cupo.
        exito, mensaje, meta = _importar(comercio_id, 'inventario.xlsx', _xlsx(reemplazo))
        _ok(bool(exito), f'.xlsx de reemplazo no bloqueado por el límite ({mensaje[:70]})')
        _ok(meta.get('bajas') == 3, f'.xlsx elimina los 3 ausentes (bajas={meta.get("bajas")})')
        _ok(meta.get('insertados') == 3, f'.xlsx inserta los 3 nuevos ({meta.get("insertados")})')
        codigos = _codigos(comercio_id)
        _ok(
            codigos == {p[0] for p in reemplazo},
            f'el inventario quedó igual al .xlsx ({sorted(codigos)})',
        )
    except Exception as exc:  # noqa: BLE001
        _ok(False, f'excepción no controlada: {type(exc).__name__}: {exc}')
    finally:
        _limpiar(get_db_connection)

    print('\n=== RESULTADO ===')
    if _ERRORES:
        for item in _ERRORES:
            print(f'  - {item}')
        print(f'FALLOS: {len(_ERRORES)}')
        return 1
    print('OK importación .csv/.xlsx + sincronización espejo (elimina ausentes y libera cupo)')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
